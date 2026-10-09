import io
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend import main


class ApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.original_db = main.DB
        main.DB = Path(self.directory.name) / 'test.sqlite3'
        self.client = TestClient(main.app)

    def tearDown(self):
        main.DB = self.original_db
        self.directory.cleanup()

    def test_import_review_history_export(self):
        csv = 'data;descricao;valor\n01/10/2026;=1+1;100,00\n02/10/2026;Tarifa;-10,00'
        internal = 'data;descricao;valor\n01/10/2026;=1+1;100,00'
        response = self.client.post('/api/sessions', data={'title': 'Outubro', 'account': 'Conta teste'}, files={'bank': ('bank.csv', csv), 'internal': ('internal.csv', internal)})
        self.assertEqual(response.status_code, 201, response.text)
        session = response.json()
        self.assertEqual([row['status'] for row in session['rows']], ['matched', 'bank_only'])
        self.assertEqual(len(self.client.get('/api/sessions').json()), 1)
        path = '/api/sessions/' + session['id']
        self.assertEqual(self.client.patch(path + '/rows/1', json={'note': 'Conferir tarifa'}).status_code, 200)
        self.assertEqual(self.client.get(path).json()['rows'][1]['note'], 'Conferir tarifa')
        exported = self.client.get(path + '/export')
        self.assertEqual(exported.status_code, 200)
        book = load_workbook(io.BytesIO(exported.content))
        self.assertEqual(book.sheetnames, ['Resumo', 'Conciliados', 'Só no banco', 'Só no interno'])
        self.assertEqual(book['Resumo']['A8'].value, 'Conciliados')
        self.assertEqual(book['Resumo']['B8'].value, 1)
        self.assertEqual(book['Resumo']['E11'].value, -0.1)
        self.assertEqual(book['Conciliados']['B2'].data_type, 's')
        self.assertEqual(book['Só no banco']['E2'].value, 'Conferir tarifa')

    def test_invalid_import_is_atomic(self):
        csv = 'data;descricao;valor\n31/02/2026;Teste;100,00'
        response = self.client.post('/api/sessions', data={'title': 'Teste', 'account': 'Teste'}, files={'bank': ('bank.csv', csv), 'internal': ('internal.csv', csv)})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get('/api/sessions').json(), [])
        self.assertEqual(self.client.get('/api/sessions/unknown').status_code, 404)

    def test_imports_two_generic_files_in_either_order(self):
        bank = 'data;descricao;valor;saldo\n01/10/2026;Pix;100,00;100,00'
        internal = 'data;descricao;valor;categoria\n01/10/2026;Pix;100,00;Receita'
        response = self.client.post('/api/sessions', data={'title': 'Teste', 'account': 'Conta teste'}, files=[('files', ('controle.csv', internal)), ('files', ('extrato.csv', bank))])
        self.assertEqual(response.status_code, 201, response.text)
        self.assertEqual(response.json()['rows'][0]['status'], 'matched')


if __name__ == '__main__':
    unittest.main()
