"""Strict import and one-to-one reconciliation; money is stored in cents."""
import csv
import io
import unicodedata
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation


def cents(value: str) -> int:
    value = str(value).strip().replace('R$', '').replace(' ', '')
    if ',' in value:
        value = value.replace('.', '').replace(',', '.')
    try:
        amount = Decimal(value)
        if not amount.is_finite() or amount != amount.quantize(Decimal('.01')):
            raise ValueError('O valor deve ter no máximo duas casas decimais.')
        if abs(amount) > Decimal('999999999999.99'):
            raise ValueError('Valor fora do limite permitido.')
        return int(amount * 100)
    except InvalidOperation as exc:
        raise ValueError('Valor monetário inválido.') from exc


def date(value: str) -> str:
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(value.strip(), fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError('Data inválida. Use DD/MM/AAAA ou AAAA-MM-DD.')


def normalize(value: str) -> str:
    return ''.join(c for c in unicodedata.normalize('NFD', value.lower().strip()) if unicodedata.category(c) != 'Mn')


ALIASES = {
    'data': 'date', 'date': 'date', 'dt': 'date', 'data lancamento': 'date', 'data movimento': 'date',
    'data pagamento': 'date', 'data emissao': 'date', 'posting date': 'date', 'transaction date': 'date',
    'descricao': 'description', 'descricao do lancamento': 'description', 'historico': 'description',
    'historico do lancamento': 'description', 'description': 'description', 'memo': 'description',
    'lancamento': 'description', 'detalhe': 'description', 'favorecido': 'description', 'beneficiario': 'description',
    'estabelecimento': 'description', 'complemento': 'description',
    'valor': 'amount', 'amount': 'amount', 'value': 'amount', 'valor lancamento': 'amount',
    'valor total': 'amount', 'valor liquido': 'amount', 'transaction amount': 'amount',
}
BANK_HINTS = {'saldo', 'banco', 'agencia', 'conta corrente', 'codigo banco', 'numero documento'}
INTERNAL_HINTS = {'centro de custo', 'conta contabil', 'categoria', 'fornecedor', 'cliente', 'plano de contas', 'sistema'}


def _decode(content: bytes) -> str:
    try:
        return content.decode('utf-8-sig')
    except UnicodeDecodeError:
        return content.decode('cp1252')


def _records_to_rows(records: list[dict]) -> list[dict]:
    if not records:
        raise ValueError('O arquivo não contém lançamentos.')
    headers = list(records[0].keys())
    mapping = {str(key): ALIASES.get(normalize(str(key))) for key in headers}
    mapped = [value for value in mapping.values() if value]
    if not {'date', 'description', 'amount'}.issubset(set(mapped)):
        raise ValueError('Colunas necessárias: data, descrição/histórico e valor.')
    columns = {target: [key for key, value in mapping.items() if value == target] for target in ('date', 'description', 'amount')}
    preferred = {'date': ('data', 'date', 'dt'), 'description': ('descricao', 'description', 'historico', 'memo', 'lancamento'), 'amount': ('valor', 'amount', 'value')}
    for target, keys in columns.items():
        keys.sort(key=lambda key: preferred[target].index(normalize(key)) if normalize(key) in preferred[target] else len(preferred[target]))
    rows = []
    for number, record in enumerate(records, 2):
        try:
            if None in record:
                raise ValueError('Quantidade de colunas invalida.')
            values = {target: record.get(keys[0]) for target, keys in columns.items()}
            description = str(values['description'] or '').strip()
            if not description or len(description) > 500:
                raise ValueError('Descrição vazia ou maior que 500 caracteres.')
            raw_date = values['date']
            transaction_date = raw_date.strftime('%Y-%m-%d') if hasattr(raw_date, 'strftime') else date(str(raw_date))
            rows.append({'date': transaction_date, 'description': description, 'amount': cents(str(values['amount']))})
        except (ValueError, AttributeError, TypeError) as exc:
            raise ValueError(f'Linha {number}: {exc}') from exc
    return rows


def _csv_records(text: str) -> tuple[list[dict], list[str]]:
    first_line = text.splitlines()[0] if text.splitlines() else ''
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=';,\t|')
    except csv.Error:
        delimiter = max((';', ',', '\t', '|'), key=first_line.count)
        if first_line.count(delimiter) == 0:
            raise ValueError('Arquivo delimitado inválido. Verifique o cabeçalho e as colunas.')
        dialect = csv.excel
        dialect.delimiter = delimiter
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise ValueError('Arquivo sem cabeçalho.')
    return list(reader), reader.fieldnames


def _detect_role(filename: str, headers: list[str], is_ofx: bool = False) -> str | None:
    if is_ofx:
        return 'bank'
    normalized = {normalize(str(header)) for header in headers}
    bank_score = sum(hint in normalized for hint in BANK_HINTS)
    internal_score = sum(hint in normalized for hint in INTERNAL_HINTS)
    if bank_score != internal_score:
        return 'bank' if bank_score > internal_score else 'internal'
    name = normalize(filename)
    bank_name = any(term in name for term in ('extrato', 'banco', 'bank'))
    internal_name = any(term in name for term in ('interno', 'sistema', 'controle', 'erp'))
    return ('bank' if bank_name else 'internal') if bank_name != internal_name else None


def parse_upload(content: bytes, filename: str) -> tuple[str | None, list[dict]]:
    extension = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''
    if extension in {'csv', 'tsv', 'txt'}:
        records, headers = _csv_records(_decode(content))
        rows = _records_to_rows(records)
        role = _detect_role(filename, headers)
    elif extension == 'ofx':
        from ofxparse import OfxParser
        try:
            ofx = OfxParser.parse(io.StringIO(_decode(content)))
            if len(ofx.accounts) != 1:
                raise ValueError('Importe um OFX de uma única conta.')
            rows = [{'date': t.date.date().isoformat(), 'description': t.memo or t.payee or 'Sem descrição', 'amount': cents(str(t.amount))} for t in ofx.account.statement.transactions]
        except Exception as exc:
            raise ValueError('OFX inválido ou com múltiplas contas. Exporte uma única conta.') from exc
        role = 'bank'
    elif extension == 'xlsx':
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(value or '').strip() for value in next(values, ())]
            records = [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]
            rows = _records_to_rows(records)
            role = _detect_role(filename, headers)
            workbook.close()
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Planilha XLSX inválida ou protegida por senha.') from exc
    else:
        raise ValueError('Formato não suportado. Envie CSV, TSV, TXT delimitado, XLSX ou OFX.')
    if not rows or len(rows) > 10000:
        raise ValueError('O arquivo deve conter entre 1 e 10.000 transações.')
    return role, rows


def parse_file(content: bytes, filename: str) -> list[dict]:
    return parse_upload(content, filename)[1]


def reconcile(bank: list[dict], internal: list[dict]) -> list[dict]:
    candidates = defaultdict(list)
    for index, row in enumerate(internal):
        candidates[(row['date'], row['amount'])].append(index)
    used = set()
    result = []
    for row in bank:
        choices = candidates[(row['date'], row['amount'])]
        available = [i for i in choices if i not in used]
        exact = [i for i in available if normalize(internal[i]['description']) == normalize(row['description'])]
        # Ambiguous equal-value entries remain pending instead of arbitrary matching.
        match = exact[0] if len(exact) == 1 else (available[0] if len(available) == 1 else None)
        if match is not None:
            used.add(match)
        result.append({'bank': row, 'internal': internal[match] if match is not None else None, 'status': 'matched' if match is not None else 'bank_only', 'note': ''})
    result.extend({'bank': None, 'internal': row, 'status': 'internal_only', 'note': ''} for i, row in enumerate(internal) if i not in used)
    return result
