import io
import json
import os
import sqlite3
from contextlib import closing, contextmanager
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

from backend.domain import parse_upload, reconcile

app = FastAPI(title='Prumo · Conciliação bancária')
app.add_middleware(CORSMiddleware, allow_origins=['http://localhost:5173', 'http://127.0.0.1:5173'], allow_methods=['GET', 'POST', 'PATCH'], allow_headers=['Content-Type'])
DB = Path(os.environ.get('PRUMO_DB', str(Path(__file__).with_name('prumo.sqlite3'))))


@contextmanager
def database():
    with closing(sqlite3.connect(DB)) as connection, connection:
        connection.row_factory = sqlite3.Row
        connection.execute('CREATE TABLE IF NOT EXISTS sessions (id TEXT PRIMARY KEY, title TEXT, account TEXT, created TEXT, data TEXT)')
        yield connection


def load(session_id):
    with database() as connection:
        record = connection.execute('SELECT * FROM sessions WHERE id=?', (session_id,)).fetchone()
    if not record:
        raise HTTPException(404, 'Conciliação não encontrada.')
    result = dict(record)
    result['rows'] = json.loads(result.pop('data'))
    return result


@app.get('/api/health')
def health():
    return {'status': 'ok'}


@app.get('/api/sessions')
def sessions():
    with database() as connection:
        return [dict(row) for row in connection.execute('SELECT id,title,account,created FROM sessions ORDER BY created DESC')]


@app.get('/api/sessions/{session_id}')
def get_session(session_id: str):
    return load(session_id)


@app.post('/api/sessions', status_code=201)
async def create_session(title: str = Form(min_length=1, max_length=100), account: str = Form(min_length=1, max_length=100), files: list[UploadFile] = File(default=[]), bank: UploadFile | None = File(default=None), internal: UploadFile | None = File(default=None)):
    if not title.strip() or not account.strip():
        raise HTTPException(422, 'Informe o nome e a conta.')
    uploads = files or [upload for upload in (bank, internal) if upload is not None]
    if len(uploads) != 2:
        raise HTTPException(422, 'Selecione exatamente dois arquivos: o extrato bancário e os lançamentos do sistema.')
    parsed = []
    for upload in uploads:
        content = await upload.read(5 * 1024 * 1024 + 1)
        if len(content) > 5 * 1024 * 1024:
            raise HTTPException(413, 'Cada arquivo deve ter até 5 MB.')
        try:
            parsed.append((upload.filename or 'arquivo', *parse_upload(content, upload.filename or '')))
        except ValueError as exc:
            raise HTTPException(422, f'{upload.filename}: {exc}') from exc
        finally:
            await upload.close()
    roles = [entry[1] for entry in parsed]
    if roles.count('bank') == 1 and roles.count('internal') == 0:
        roles[roles.index(None)] = 'internal'
    elif roles.count('internal') == 1 and roles.count('bank') == 0:
        roles[roles.index(None)] = 'bank'
    if set(roles) != {'bank', 'internal'}:
        raise HTTPException(422, 'Não consegui distinguir os arquivos. Use um extrato OFX ou CSV com coluna de saldo/banco e um controle com coluna de categoria, fornecedor ou centro de custo. Também é possível identificar pelo nome do arquivo (extrato/controle).')
    by_role = {role: rows for (_, _, rows), role in zip(parsed, roles)}
    rows = reconcile(by_role['bank'], by_role['internal'])
    for index, row in enumerate(rows):
        row['id'] = str(index)
    session_id = str(uuid4())
    with database() as connection:
        connection.execute('INSERT INTO sessions VALUES (?,?,?,?,?)', (session_id, title.strip(), account.strip(), datetime.now(timezone.utc).isoformat(), json.dumps(rows)))
    return load(session_id)


class Review(BaseModel):
    note: str = Field(max_length=1000)


@app.patch('/api/sessions/{session_id}/rows/{row_id}')
def review(session_id: str, row_id: str, body: Review):
    # Read and write under one transaction to avoid losing simultaneous notes.
    with database() as connection:
        connection.execute('BEGIN IMMEDIATE')
        record = connection.execute('SELECT data FROM sessions WHERE id=?', (session_id,)).fetchone()
        if not record:
            raise HTTPException(404, 'Conciliação não encontrada.')
        rows = json.loads(record['data'])
        row = next((r for r in rows if r['id'] == row_id), None)
        if row is None:
            raise HTTPException(404, 'Transação não encontrada.')
        row['note'] = body.note.strip()
        connection.execute('UPDATE sessions SET data=? WHERE id=?', (json.dumps(rows), session_id))
    return load(session_id)


@app.get('/api/sessions/{session_id}/export')
def export(session_id: str):
    session = load(session_id)
    book = Workbook()
    book.remove(book.active)
    summary = book.create_sheet('Resumo')
    rows = session['rows']
    transactions = [transaction for row in rows for transaction in (row['bank'], row['internal']) if transaction]
    dates = [transaction['date'] for transaction in transactions]
    summary.append(['Relatório de conciliação'])
    summary.append(['Conciliação', session['title']])
    summary.append(['Conta', session['account']])
    summary.append(['Criado em', session['created']])
    summary.append(['Período', f'{min(dates)} a {max(dates)}' if dates else 'Sem transações'])
    summary.append([])
    summary.append(['Situação', 'Lançamentos', 'Total banco (R$)', 'Total interno (R$)', 'Diferença (R$)'])
    for status, label in [('matched', 'Conciliados'), ('bank_only', 'Pendentes no banco'), ('internal_only', 'Pendentes no controle')]:
        group = [row for row in rows if row['status'] == status]
        bank_total = sum((row['bank'] or {}).get('amount', 0) for row in group)
        internal_total = sum((row['internal'] or {}).get('amount', 0) for row in group)
        summary.append([label, len(group), bank_total / 100, internal_total / 100, (bank_total - internal_total) / 100])
    bank_total = sum((row['bank'] or {}).get('amount', 0) for row in rows)
    internal_total = sum((row['internal'] or {}).get('amount', 0) for row in rows)
    summary.append(['Total', len(rows), bank_total / 100, internal_total / 100, (bank_total - internal_total) / 100])
    for cell in summary[1]:
        cell.font = Font(color='FFFFFF', bold=True, size=14)
        cell.fill = PatternFill('solid', fgColor='183E35')
    for cell in summary[7]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='183E35')
    for row in summary.iter_rows(min_row=8, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00'
    for row in summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    summary.column_dimensions['A'].width = 30
    summary.column_dimensions['B'].width = 22
    for column in ('C', 'D', 'E'):
        summary.column_dimensions[column].width = 24
    summary.freeze_panes = 'A8'
    summary.auto_filter.ref = 'A7:E10'
    for status, title in [('matched', 'Conciliados'), ('bank_only', 'Só no banco'), ('internal_only', 'Só no interno')]:
        sheet = book.create_sheet(title)
        sheet.append(['Data', 'Descrição banco', 'Descrição interna', 'Valor (R$)', 'Observação'])
        for row in session['rows']:
            if row['status'] != status:
                continue
            transaction = row['bank'] or row['internal']
            sheet.append([transaction['date'], (row['bank'] or {}).get('description', ''), (row['internal'] or {}).get('description', ''), transaction['amount'] / 100, row['note']])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Untrusted text must never become an Excel formula.
            sheet.cell(sheet.max_row, 4).number_format = '#,##0.00'
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='183E35')
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in [('A', 14), ('B', 40), ('C', 40), ('D', 20), ('E', 60)]:
            sheet.column_dimensions[column].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename="conciliacao.xlsx"'})
